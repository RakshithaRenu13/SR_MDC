import os
import sqlite3
from io import BytesIO
from datetime import datetime

import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
)

import streamlit as st


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="Eaton MDC Solution Configurator",
    page_icon="⚡",
    layout="wide",
    initial_sidebar_state="expanded",
)


# ============================================================
# GLOBAL CSS
# ============================================================

st.markdown(
    """
    <style>

    .stApp {
        background:#F8FAFC;
    }

    .block-container {
        padding-top:0.65rem;
        padding-bottom:1rem;
        padding-left:1.2rem;
        padding-right:1.2rem;
        max-width:100%;
    }

    .mdc-title-card {
        background:linear-gradient(
            90deg,
            #075EA8 0%,
            #003B71 100%
        );
        color:white;
        border-radius:7px;
        padding:10px 16px 9px 16px;
        margin-bottom:9px;
        box-shadow:0 2px 6px rgba(0,0,0,0.08);
    }

    .mdc-title {
        font-size:22px;
        font-weight:700;
        line-height:1.2;
    }

    .mdc-subtitle {
        font-size:12px;
        margin-top:3px;
        opacity:0.92;
    }

    .mdc-card-heading {
        display:flex;
        align-items:center;
        gap:8px;
        color:#003B71;
        font-size:13px;
        font-weight:700;
        margin-bottom:7px;
        text-transform:uppercase;
    }

    .mdc-number {
        background:#007AC2;
        color:white;
        border-radius:4px;
        padding:3px 7px;
        font-size:11px;
        font-weight:700;
    }

    .mdc-mini-heading {
        color:#005EB8;
        font-size:11px;
        font-weight:700;
        margin:5px 0 3px 0;
    }

    div[data-testid="stVerticalBlockBorderWrapper"] {
        border-color:#D9E6F0;
        border-radius:7px;
    }

    .stButton button,
    .stDownloadButton button {
        border-radius:5px;
    }

    div[data-testid="stMetric"] {
        background:white;
        border:1px solid #D9E6F0;
        border-radius:6px;
        padding:5px 9px;
    }

    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# PATHS
# ============================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

MASTER_FILE = os.path.join(
    BASE_DIR,
    "MDC_Master_V1.xlsx"
)

TRACKING_DB = os.path.join(
    BASE_DIR,
    "MDC_Tracking.db"
)

DEMO_INTERNAL_PASSWORD = "MDC@123"


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def clean_text(value):
    if pd.isna(value):
        return ""

    return str(value).strip()


def numeric(value):
    try:
        if pd.isna(value):
            return 0.0

        text = str(value).replace(",", "").strip()

        if text == "":
            return 0.0

        return float(text)

    except Exception:
        return 0.0


def money(value):
    try:
        return f"₹{numeric(value):,.2f}"
    except Exception:
        return "₹0.00"


def internal_password():
    return DEMO_INTERNAL_PASSWORD


# ============================================================
# DATABASE
# ============================================================

def init_tracking_db():

    conn = sqlite3.connect(TRACKING_DB)
    cur = conn.cursor()

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS configurations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            configuration_id TEXT,
            customer_name TEXT,
            customer_place TEXT,
            problem TEXT,
            solution TEXT,
            mdc_type TEXT,
            configuration TEXT,
            base_cost REAL,
            optional_cost REAL,
            pdu_cost REAL,
            total_cost REAL,
            margin_pct REAL,
            freight REAL,
            installation REAL,
            warranty_pct REAL,
            margin_price REAL,
            final_selling_price REAL,
            warranty_amount REAL,
            created_at TEXT,
            user_code TEXT
        )
        """
    )

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS configuration_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            configuration_id TEXT,
            component_type TEXT,
            part_code TEXT,
            description TEXT,
            quantity REAL,
            uom TEXT,
            unit_cost REAL,
            total_cost REAL,
            unit_price REAL,
            total_price REAL
        )
        """
    )

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS download_counter (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            download_count INTEGER DEFAULT 0
        )
        """
    )

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS session_counter (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_count INTEGER DEFAULT 0
        )
        """
    )

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS user_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_code TEXT,
            customer_name TEXT,
            created_at TEXT
        )
        """
    )

    columns = [
        x[1]
        for x in cur.execute(
            "PRAGMA table_info(configurations)"
        ).fetchall()
    ]

    if "user_code" not in columns:
        cur.execute(
            "ALTER TABLE configurations ADD COLUMN user_code TEXT"
        )

    conn.commit()

    # Ensure counter rows exist
    cur.execute(
        "SELECT COUNT(*) FROM download_counter"
    )

    if cur.fetchone()[0] == 0:
        cur.execute(
            "INSERT INTO download_counter(download_count) VALUES(0)"
        )

    cur.execute(
        "SELECT COUNT(*) FROM session_counter"
    )

    if cur.fetchone()[0] == 0:
        cur.execute(
            "INSERT INTO session_counter(user_count) VALUES(0)"
        )

    conn.commit()
    conn.close()


init_tracking_db()


# ============================================================
# COUNTERS
# ============================================================

def increment_user_count():

    conn = sqlite3.connect(TRACKING_DB)
    cur = conn.cursor()

    cur.execute(
        """
        UPDATE session_counter
        SET user_count = user_count + 1
        WHERE id = 1
        """
    )

    cur.execute(
        """
        SELECT user_count
        FROM session_counter
        WHERE id = 1
        """
    )

    value = cur.fetchone()[0]

    conn.commit()
    conn.close()

    return value


def get_user_count():

    conn = sqlite3.connect(TRACKING_DB)

    value = conn.execute(
        """
        SELECT user_count
        FROM session_counter
        WHERE id = 1
        """
    ).fetchone()[0]

    conn.close()

    return value


def increment_download_count():

    conn = sqlite3.connect(TRACKING_DB)
    cur = conn.cursor()

    cur.execute(
        """
        UPDATE download_counter
        SET download_count = download_count + 1
        WHERE id = 1
        """
    )

    conn.commit()
    conn.close()


# ============================================================
# CONFIGURATION ID
# ============================================================

def generate_configuration_id():

    now = datetime.now()

    conn = sqlite3.connect(TRACKING_DB)

    count = conn.execute(
        """
        SELECT COUNT(*)
        FROM configurations
        WHERE DATE(created_at) = DATE(?)
        """,
        (now.isoformat(),),
    ).fetchone()[0]

    conn.close()

    return (
        f"MDC-{now.strftime('%Y%m%d')}-"
        f"{count + 1:04d}"
    )


# ============================================================
# MASTER EXCEL
# ============================================================

@st.cache_data
def load_master():

    if not os.path.exists(MASTER_FILE):
        st.error(
            "MDC_Master_V1.xlsx was not found "
            "in the application folder."
        )
        st.stop()

    xl = pd.ExcelFile(MASTER_FILE)

    first_sheet = xl.sheet_names[0]

    raw = pd.read_excel(
        MASTER_FILE,
        sheet_name=first_sheet,
        header=None
    )

    configs = []
    components = []

    # --------------------------------------------------------
    # SINGLE RACK CONFIGURATIONS
    # Existing architecture preserved
    # --------------------------------------------------------

    single_blocks = [
        ("Configuration 1", 0, 25, 0, 5),
        ("Configuration 3", 0, 25, 5, 10),
        ("Configuration 2", 25, 50, 0, 5),
        ("Configuration 4", 25, 50, 5, 10),
    ]

    for config_name, r1, r2, c1, c2 in single_blocks:

        block = raw.iloc[
            r1:r2,
            c1:c2
        ].copy()

        if block.empty:
            continue

        headers = [
            clean_text(x)
            for x in block.iloc[0].tolist()
        ]

        if len(headers) < 5:
            continue

        data = block.iloc[1:].copy()
        data.columns = headers

        code_col = headers[0]
        desc_col = headers[1]
        qty_col = headers[2]
        uom_col = headers[3]
        cost_col = headers[4]

        config_cost = 0.0

        for _, row in data.iterrows():

            part_code = clean_text(
                row.get(code_col, "")
            )

            description = clean_text(
                row.get(desc_col, "")
            )

            quantity = numeric(
                row.get(qty_col, 0)
            )

            uom = clean_text(
                row.get(uom_col, "")
            )

            unit_cost = numeric(
                row.get(cost_col, 0)
            )

            if not part_code and not description:
                continue

            total_cost = quantity * unit_cost

            config_cost += total_cost

            components.append(
                {
                    "MDC Type": "Single Rack",
                    "Configuration": config_name,
                    "Part Code": part_code,
                    "Description": description,
                    "Quantity": quantity,
                    "UOM": uom,
                    "Unit Cost": unit_cost,
                }
            )

        configs.append(
            {
                "MDC Type": "Single Rack",
                "Configuration": config_name,
                "Configuration Title": config_name,
                "Base Cost": config_cost,
            }
        )

    # --------------------------------------------------------
    # STANDARD MULTIRACK CONFIGURATIONS
    #
    # IMPORTANT:
    # No rack-size customization is introduced here.
    #
    # Multirack is selected only as a STANDARD configuration.
    # --------------------------------------------------------

    for n in range(1, 10):

        configs.append(
            {
                "MDC Type": "Multirack",
                "Configuration":
                    f"Configuration {n}",
                "Configuration Title":
                    f"Multirack Configuration {n}",
                "Base Cost": 0.0,
            }
        )

    # --------------------------------------------------------
    # SEARCH FOR OPTIONAL ITEMS
    # --------------------------------------------------------

    accessories = []

    optional_start = None

    for r in range(len(raw)):

        row_text = " ".join(
            clean_text(v).upper()
            for v in raw.iloc[r].tolist()
        )

        if "OTHER OPTIONAL ITEMS" in row_text:
            optional_start = r
            break

    if optional_start is not None:

        for r in range(
            optional_start + 1,
            len(raw)
        ):

            row = raw.iloc[r]

            values = [
                clean_text(v)
                for v in row.tolist()
            ]

            if not any(values):
                continue

            part_code = values[0] if len(values) > 0 else ""
            description = values[1] if len(values) > 1 else ""
            uom = values[2] if len(values) > 2 else ""
            unit_cost = (
                numeric(values[3])
                if len(values) > 3
                else 0.0
            )

            if not part_code and not description:
                continue

            accessories.append(
                {
                    "Part Code": part_code,
                    "Description": description,
                    "UOM": uom,
                    "Unit Cost": unit_cost,
                }
            )

    accessories_df = pd.DataFrame(
        accessories,
        columns=[
            "Part Code",
            "Description",
            "UOM",
            "Unit Cost",
        ],
    )

    # --------------------------------------------------------
    # SEARCH FOR SINGLE PHASE PDU
    # --------------------------------------------------------

    pdus = []

    pdu_start = None

    for r in range(len(raw)):

        row_text = " ".join(
            clean_text(v).upper()
            for v in raw.iloc[r].tolist()
        )

        if "SINGLE PHASE PDU" in row_text:
            pdu_start = r
            break

    if pdu_start is not None:

        for r in range(
            pdu_start + 1,
            len(raw)
        ):

            row = raw.iloc[r]

            values = [
                clean_text(v)
                for v in row.tolist()
            ]

            if not any(values):
                continue

            part_code = values[0] if len(values) > 0 else ""
            description = values[1] if len(values) > 1 else ""
            c13 = values[2] if len(values) > 2 else ""
            c19 = values[3] if len(values) > 3 else ""
            pdu_type = values[4] if len(values) > 4 else ""
            uom = values[5] if len(values) > 5 else ""
            unit_cost = (
                numeric(values[6])
                if len(values) > 6
                else 0.0
            )

            if not part_code and not description:
                continue

            pdus.append(
                {
                    "Part Code": part_code,
                    "Description": description,
                    "C13": c13,
                    "C19": c19,
                    "Type": pdu_type,
                    "UOM": uom,
                    "Unit Cost": unit_cost,
                }
            )

    pdus_df = pd.DataFrame(
        pdus,
        columns=[
            "Part Code",
            "Description",
            "C13",
            "C19",
            "Type",
            "UOM",
            "Unit Cost",
        ],
    )

    configs_df = pd.DataFrame(configs)

    components_df = pd.DataFrame(
        components,
        columns=[
            "MDC Type",
            "Configuration",
            "Part Code",
            "Description",
            "Quantity",
            "UOM",
            "Unit Cost",
        ],
    )

    return (
        configs_df,
        components_df,
        accessories_df,
        pdus_df,
    )


configs_df, components_df, accessories_df, pdus_df = load_master()


# ============================================================
# SESSION DEFAULTS
# ============================================================

defaults = {
    "mode": "Sales",
    "authenticated": False,

    "customer_name": "",
    "customer_place": "",
    "problem": "",
    "solution": "",

    "mdc_type": "Single Rack",
    "configuration": "Configuration 1",

    "accessory_qty": {},
    "pdu_qty": {},

    "pdu_type_selection": "None",
    "pdu_model_selection": None,

    "fire_suppression_selection": "None",
    "camera_quantity": 1,

    "margin_pct": 20.0,
    "freight": 0.0,
    "installation": 0.0,
    "warranty_pct": 0.0,

    "configuration_id": None,
    "configuration_saved": False,

    "user_code": "—",
    "user_count": 0,
    "session_counted": False,
}


for key, value in defaults.items():

    if key not in st.session_state:
        st.session_state[key] = value


if st.session_state.configuration_id is None:

    st.session_state.configuration_id = (
        generate_configuration_id()
    )


if not st.session_state.session_counted:

    st.session_state.user_count = (
        increment_user_count()
    )

    st.session_state.session_counted = True

else:

    st.session_state.user_count = (
        get_user_count()
    )


# ============================================================
# UI HELPERS
# ============================================================

def section_header(text):

    st.markdown(
        f"""
        <div style="
            background:#003B71;
            color:white;
            padding:6px 10px;
            border-radius:5px;
            margin:9px 0 7px 0;
            font-size:13px;
            font-weight:700;
        ">
            {text}
        </div>
        """,
        unsafe_allow_html=True,
    )


def price_box(label, value):

    st.markdown(
        f"""
        <div style="
            padding:4px 0 12px 0;
            min-height:82px;
            overflow:visible;
        ">
            <div style="
                font-size:11px;
                color:#64748B;
                font-weight:600;
                margin-bottom:4px;
            ">
                {label}
            </div>

            <div style="
                font-size:16px;
                color:#003B71;
                font-weight:700;
            ">
                {money(value)}
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


# ============================================================
# USER CODE
# ============================================================

def generate_user_code():

    type_code = (
        "SR"
        if st.session_state.mdc_type == "Single Rack"
        else "MR"
    )

    config_text = (
        st.session_state.configuration
    )

    config_number = "1"

    for char in config_text:

        if char.isdigit():
            config_number = char
            break

    config_code = f"C{config_number}"

    signature_source = []

    # PDU
    pdu_part = next(
        iter(
            st.session_state.pdu_qty.keys()
        ),
        "",
    )

    if pdu_part:
        signature_source.append(
            pdu_part
        )

    # Accessories
    for part in sorted(
        st.session_state.accessory_qty.keys()
    ):

        if numeric(
            st.session_state.accessory_qty.get(
                part,
                0
            )
        ) > 0:

            signature_source.append(part)

    signature = ""

    for value in signature_source:

        for char in str(value):

            if char.isalnum():
                signature += char.upper()

            if len(signature) >= 4:
                break

        if len(signature) >= 4:
            break

    if len(signature) < 4:

        signature = (
            signature + "0000"
        )[:4]

    return (
        f"{type_code}-"
        f"{config_code}-"
        f"{signature}-0001"
    )


def handle_excel_download():

    st.session_state.user_code = (
        generate_user_code()
    )

    conn = sqlite3.connect(TRACKING_DB)

    conn.execute(
        """
        INSERT INTO user_history(
            user_code,
            customer_name,
            created_at
        )
        VALUES (?, ?, ?)
        """,
        (
            st.session_state.user_code,
            st.session_state.customer_name,
            datetime.now().isoformat(),
        ),
    )

    conn.commit()
    conn.close()


def render_user_info_panel(container):

    with container:

        st.markdown(
            f"""
            <div style="
                display:flex;
                justify-content:space-between;
                align-items:center;
                gap:10px;
                background:white;
                border:1px solid #D9E6F0;
                border-radius:6px;
                padding:7px 10px;
                margin-bottom:7px;
                font-size:11px;
            ">

                <div>
                    <b>User Code</b><br>
                    <span style="color:#003B71;">
                        {st.session_state.user_code}
                    </span>
                </div>

                <div>
                    <b>User Count</b><br>
                    {st.session_state.user_count}
                </div>

                <div>
                    <b>Date</b><br>
                    {datetime.now().strftime("%d-%m-%Y")}
                </div>

                <div>
                    <b>Access</b><br>
                    {"Internal" if st.session_state.authenticated else "Sales"}
                </div>

            </div>
            """,
            unsafe_allow_html=True,
        )


# ============================================================
# DATA LOOKUPS
# ============================================================

def selected_config_record():

    matches = configs_df[
        (configs_df["MDC Type"] == st.session_state.mdc_type)
        &
        (
            configs_df["Configuration"]
            == st.session_state.configuration
        )
    ]

    if matches.empty:
        return None

    return matches.iloc[0].to_dict()


def selected_components():

    matches = components_df[
        (components_df["MDC Type"] == st.session_state.mdc_type)
        &
        (
            components_df["Configuration"]
            == st.session_state.configuration
        )
    ].copy()

    return matches


# ============================================================
# BUILD BOM
# ============================================================

def build_bom():

    rows = []

    # --------------------------------------------------------
    # MAIN MDC
    # --------------------------------------------------------

    config_record = selected_config_record()

    if config_record is not None:

        rows.append(
            {
                "S.No.": 1,
                "Component Type": "Main MDC",
                "Part Code": "801029209",
                "Description":
                    (
                        "SINGLE RACK MDC"
                        if st.session_state.mdc_type
                        == "Single Rack"
                        else "MULTIRACK MDC"
                    ),
                "Quantity": 1,
                "UOM": "EA",
                "Unit Cost": numeric(
                    config_record.get(
                        "Base Cost",
                        0
                    )
                ),
                "Total Cost": numeric(
                    config_record.get(
                        "Base Cost",
                        0
                    )
                ),
                "Source": "Base",
            }
        )

    # --------------------------------------------------------
    # CONFIGURATION COMPONENTS
    # --------------------------------------------------------

    selected = selected_components()

    for _, r in selected.iterrows():

        part_code = clean_text(
            r["Part Code"]
        )

        description = clean_text(
            r["Description"]
        )

        quantity = numeric(
            r["Quantity"]
        )

        uom = clean_text(
            r["UOM"]
        )

        unit_cost = numeric(
            r["Unit Cost"]
        )

        if not part_code and not description:
            continue

        rows.append(
            {
                "S.No.": "",
                "Component Type":
                    "Base (Configuration)",
                "Part Code": part_code,
                "Description": description,
                "Quantity": quantity,
                "UOM": uom,
                "Unit Cost": unit_cost,
                "Total Cost":
                    quantity * unit_cost,
                "Source": "Base",
            }
        )

    # --------------------------------------------------------
    # PDU
    # --------------------------------------------------------

    for _, r in pdus_df.iterrows():

        part = clean_text(
            r["Part Code"]
        )

        qty = numeric(
            st.session_state.pdu_qty.get(
                part,
                0
            )
        )

        if qty <= 0:
            continue

        description = clean_text(
            r["Description"]
        )

        pdu_type = clean_text(
            r["Type"]
        )

        c13 = clean_text(
            r["C13"]
        )

        c19 = clean_text(
            r["C19"]
        )

        details = []

        if pdu_type:
            details.append(pdu_type)

        if c13:
            details.append(
                f"C13: {c13}"
            )

        if c19:
            details.append(
                f"C19: {c19}"
            )

        if details:
            description += (
                " — "
                + " | ".join(details)
            )

        unit_cost = numeric(
            r["Unit Cost"]
        )

        rows.append(
            {
                "S.No.": "",
                "Component Type": "PDU",
                "Part Code": part,
                "Description": description,
                "Quantity": qty,
                "UOM": clean_text(r["UOM"]),
                "Unit Cost": unit_cost,
                "Total Cost":
                    qty * unit_cost,
                "Source": "PDU",
            }
        )

    # --------------------------------------------------------
    # OPTIONAL ACCESSORIES
    # --------------------------------------------------------

    for _, r in accessories_df.iterrows():

        part = clean_text(
            r["Part Code"]
        )

        qty = numeric(
            st.session_state.accessory_qty.get(
                part,
                0
            )
        )

        if qty <= 0:
            continue

        unit_cost = numeric(
            r["Unit Cost"]
        )

        rows.append(
            {
                "S.No.": "",
                "Component Type":
                    "Optional Accessory",
                "Part Code": part,
                "Description":
                    clean_text(
                        r["Description"]
                    ),
                "Quantity": qty,
                "UOM":
                    clean_text(
                        r["UOM"]
                    ),
                "Unit Cost": unit_cost,
                "Total Cost":
                    qty * unit_cost,
                "Source": "Optional",
            }
        )

    return pd.DataFrame(rows)


# ============================================================
# COST SUMMARY
# ============================================================

def cost_summary(bom):

    if bom.empty:
        return 0.0, 0.0, 0.0, 0.0

    base_cost = numeric(
        bom.loc[
            bom["Source"] == "Base",
            "Total Cost"
        ].sum()
    )

    optional_cost = numeric(
        bom.loc[
            bom["Source"] == "Optional",
            "Total Cost"
        ].sum()
    )

    pdu_cost = numeric(
        bom.loc[
            bom["Source"] == "PDU",
            "Total Cost"
        ].sum()
    )

    total_cost = (
        base_cost
        + optional_cost
        + pdu_cost
    )

    return (
        base_cost,
        optional_cost,
        pdu_cost,
        total_cost,
    )


# ============================================================
# SELLING PRICE
# ============================================================

def add_selling_prices(
    bom,
    total_cost,
    margin_pct,
    freight,
    installation,
):

    output = bom.copy()

    output["Unit Price"] = (
        output["Unit Cost"]
    )

    output["Total Price"] = (
        output["Unit Price"]
        * output["Quantity"]
    )

    if margin_pct < 100:

        margin_price = (
            total_cost
            / (
                1
                - margin_pct / 100
            )
        )

    else:

        margin_price = 0.0

    final_selling_price = (
        margin_price
        + freight
        + installation
    )

    return (
        output,
        margin_price,
        final_selling_price,
    )


# ============================================================
# EXCEL OUTPUT
# ============================================================

def excel_bytes(
    internal=False,
    bom=None,
    final_price=0.0,
    cost_data=None,
):

    output = BytesIO()

    from openpyxl import Workbook

    wb = Workbook()

    ws = wb.active
    ws.title = "MDC BOQ"

    ws.sheet_view.showGridLines = False

    blue = "003B71"
    light_blue = "D9EAF7"
    header_gray = "F4F6F8"

    title_fill = PatternFill(
        "solid",
        fgColor=blue
    )

    header_fill = PatternFill(
        "solid",
        fgColor=header_gray
    )

    selected_fill = PatternFill(
        "solid",
        fgColor=light_blue
    )

    white_font = Font(
        color="FFFFFF",
        bold=True
    )

    bold_font = Font(
        bold=True,
        color=blue
    )

    thin = Side(
        style="thin",
        color="D9E1E8"
    )

    # --------------------------------------------------------
    # TITLE
    # --------------------------------------------------------

    ws.merge_cells(
        "A1:I1"
    )

    ws["A1"] = (
        "Eaton MDC Solution Configurator"
    )

    ws["A1"].fill = title_fill
    ws["A1"].font = Font(
        color="FFFFFF",
        bold=True,
        size=16
    )

    ws["A1"].alignment = Alignment(
        horizontal="center"
    )

    ws.merge_cells(
        "A2:I2"
    )

    ws["A2"] = (
        "Modular Data Center Solution "
        "Configuration & Pricing"
    )

    ws["A2"].alignment = Alignment(
        horizontal="center"
    )

    # --------------------------------------------------------
    # CUSTOMER
    # --------------------------------------------------------

    ws["A4"] = "Customer Name"
    ws["B4"] = (
        st.session_state.customer_name
    )

    ws["A5"] = "Customer Place"
    ws["B5"] = (
        st.session_state.customer_place
    )

    ws["A6"] = "Problem Description"
    ws["B6"] = (
        st.session_state.problem
    )

    ws["A7"] = "Solution"
    ws["B7"] = (
        st.session_state.solution
    )

    ws["A8"] = "MDC Type"
    ws["B8"] = (
        st.session_state.mdc_type
    )

    ws["A9"] = "Configuration"
    ws["B9"] = (
        st.session_state.configuration
    )

    row = 11

    # --------------------------------------------------------
    # FINAL BOQ
    # --------------------------------------------------------

    ws.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=7
    )

    ws.cell(
        row,
        1
    ).value = "FINAL BOQ"

    ws.cell(
        row,
        1
    ).fill = title_fill

    ws.cell(
        row,
        1
    ).font = white_font

    ws.cell(
        row,
        1
    ).alignment = Alignment(
        horizontal="center"
    )

    row += 1

    if internal:

        headers = [
            "S.No.",
            "Part Code",
            "Description",
            "Quantity",
            "UOM",
            "Unit Cost",
            "Total Cost",
            "Unit Price",
            "Total Price",
        ]

    else:

        headers = [
            "S.No.",
            "Part Code",
            "Description",
            "Quantity",
            "UOM",
            "Unit Price",
            "Total Price",
        ]

    for col, header in enumerate(
        headers,
        start=1
    ):

        cell = ws.cell(
            row,
            col
        )

        cell.value = header
        cell.fill = header_fill
        cell.font = Font(
            bold=True
        )
        cell.alignment = Alignment(
            horizontal="center"
        )

    row += 1

    if bom is not None and not bom.empty:

        for _, r in bom.iterrows():

            if internal:

                values = [
                    r.get("S.No.", ""),
                    r.get("Part Code", ""),
                    r.get("Description", ""),
                    r.get("Quantity", ""),
                    r.get("UOM", ""),
                    numeric(
                        r.get("Unit Cost", 0)
                    ),
                    numeric(
                        r.get("Total Cost", 0)
                    ),
                    numeric(
                        r.get("Unit Price", 0)
                    ),
                    numeric(
                        r.get("Total Price", 0)
                    ),
                ]

            else:

                values = [
                    r.get("S.No.", ""),
                    r.get("Part Code", ""),
                    r.get("Description", ""),
                    r.get("Quantity", ""),
                    r.get("UOM", ""),
                    numeric(
                        r.get("Unit Price", 0)
                    ),
                    numeric(
                        r.get("Total Price", 0)
                    ),
                ]

            for col, value in enumerate(
                values,
                start=1
            ):

                ws.cell(
                    row,
                    col
                ).value = value

            row += 1

    # --------------------------------------------------------
    # PRICE SUMMARY
    # --------------------------------------------------------

    row += 1

    ws.cell(
        row,
        1
    ).value = "FINAL SELLING PRICE"

    ws.cell(
        row,
        1
    ).font = bold_font

    ws.cell(
        row,
        2
    ).value = final_price

    ws.cell(
        row,
        2
    ).font = bold_font

    if internal and cost_data:

        row += 2

        for label, value in cost_data:

            ws.cell(
                row,
                1
            ).value = label

            ws.cell(
                row,
                2
            ).value = value

            row += 1

    # --------------------------------------------------------
    # FORMATTING
    # --------------------------------------------------------

    for row_cells in ws.iter_rows():

        for cell in row_cells:

            cell.border = Border(
                bottom=thin
            )

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True
            )

    widths = {
        "A": 10,
        "B": 18,
        "C": 55,
        "D": 12,
        "E": 10,
        "F": 16,
        "G": 16,
        "H": 16,
        "I": 16,
    }

    for col, width in widths.items():
        ws.column_dimensions[
            col
        ].width = width

    ws.freeze_panes = "A13"

    wb.save(output)

    output.seek(0)

    return output.getvalue()


# ============================================================
# PDF OUTPUT
# ============================================================

def pdf_bytes(
    internal=False,
    bom=None,
    final_price=0.0,
):

    output = BytesIO()

    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "Title",
        parent=styles["Heading1"],
        alignment=TA_CENTER,
        fontSize=16,
        textColor=colors.HexColor(
            "#003B71"
        ),
    )

    normal_style = ParagraphStyle(
        "Normal",
        parent=styles["Normal"],
        fontSize=8,
    )

    story = []

    story.append(
        Paragraph(
            "Eaton MDC Solution Configurator",
            title_style
        )
    )

    story.append(
        Paragraph(
            "Modular Data Center Solution "
            "Configuration & Pricing",
            normal_style
        )
    )

    story.append(
        Spacer(1, 5 * mm)
    )

    customer_data = [
        [
            "Customer",
            st.session_state.customer_name,
            "MDC Type",
            st.session_state.mdc_type,
        ],
        [
            "Place",
            st.session_state.customer_place,
            "Configuration",
            st.session_state.configuration,
        ],
    ]

    customer_table = Table(
        customer_data,
        colWidths=[
            25 * mm,
            65 * mm,
            30 * mm,
            65 * mm,
        ],
    )

    customer_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (0, -1),
                    colors.HexColor(
                        "#F4F6F8"
                    ),
                ),
                (
                    "BACKGROUND",
                    (2, 0),
                    (2, -1),
                    colors.HexColor(
                        "#F4F6F8"
                    ),
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.HexColor(
                        "#D9E1E8"
                    ),
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, -1),
                    "Helvetica",
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (0, -1),
                    "Helvetica-Bold",
                ),
                (
                    "FONTNAME",
                    (2, 0),
                    (2, -1),
                    "Helvetica-Bold",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
            ]
        )
    )

    story.append(
        customer_table
    )

    story.append(
        Spacer(1, 5 * mm)
    )

    story.append(
        Paragraph(
            "FINAL BOQ",
            normal_style
        )
    )

    if internal:

        headers = [
            "S.No.",
            "Part Code",
            "Description",
            "Qty",
            "UOM",
            "Unit Cost",
            "Total Cost",
            "Unit Price",
            "Total Price",
        ]

    else:

        headers = [
            "S.No.",
            "Part Code",
            "Description",
            "Qty",
            "UOM",
            "Unit Price",
            "Total Price",
        ]

    table_data = [headers]

    if bom is not None and not bom.empty:

        for _, r in bom.iterrows():

            if internal:

                table_data.append(
                    [
                        clean_text(
                            r.get(
                                "S.No.",
                                ""
                            )
                        ),
                        clean_text(
                            r.get(
                                "Part Code",
                                ""
                            )
                        ),
                        clean_text(
                            r.get(
                                "Description",
                                ""
                            )
                        ),
                        str(
                            numeric(
                                r.get(
                                    "Quantity",
                                    0
                                )
                            )
                        ),
                        clean_text(
                            r.get(
                                "UOM",
                                ""
                            )
                        ),
                        money(
                            r.get(
                                "Unit Cost",
                                0
                            )
                        ),
                        money(
                            r.get(
                                "Total Cost",
                                0
                            )
                        ),
                        money(
                            r.get(
                                "Unit Price",
                                0
                            )
                        ),
                        money(
                            r.get(
                                "Total Price",
                                0
                            )
                        ),
                    ]
                )

            else:

                table_data.append(
                    [
                        clean_text(
                            r.get(
                                "S.No.",
                                ""
                            )
                        ),
                        clean_text(
                            r.get(
                                "Part Code",
                                ""
                            )
                        ),
                        clean_text(
                            r.get(
                                "Description",
                                ""
                            )
                        ),
                        str(
                            numeric(
                                r.get(
                                    "Quantity",
                                    0
                                )
                            )
                        ),
                        clean_text(
                            r.get(
                                "UOM",
                                ""
                            )
                        ),
                        money(
                            r.get(
                                "Unit Price",
                                0
                            )
                        ),
                        money(
                            r.get(
                                "Total Price",
                                0
                            )
                        ),
                    ]
                )

    if internal:

        widths = [
            14 * mm,
            28 * mm,
            100 * mm,
            15 * mm,
            15 * mm,
            25 * mm,
            27 * mm,
            25 * mm,
            27 * mm,
        ]

    else:

        widths = [
            15 * mm,
            30 * mm,
            115 * mm,
            18 * mm,
            18 * mm,
            30 * mm,
            32 * mm,
        ]

    boq_table = Table(
        table_data,
        colWidths=widths,
        repeatRows=1,
    )

    boq_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor(
                        "#003B71"
                    ),
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.35,
                    colors.HexColor(
                        "#D9E1E8"
                    ),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "ALIGN",
                    (0, 0),
                    (0, -1),
                    "CENTER",
                ),
                (
                    "ALIGN",
                    (3, 1),
                    (4, -1),
                    "CENTER",
                ),
                (
                    "ALIGN",
                    (-2, 1),
                    (-1, -1),
                    "RIGHT",
                ),
            ]
        )
    )

    story.append(
        boq_table
    )

    story.append(
        Spacer(1, 4 * mm)
    )

    final_table = Table(
        [
            [
                "FINAL SELLING PRICE",
                money(final_price),
            ]
        ],
        colWidths=[
            50 * mm,
            40 * mm,
        ],
    )

    final_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, -1),
                    colors.HexColor(
                        "#D9EAF7"
                    ),
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, -1),
                    colors.HexColor(
                        "#003B71"
                    ),
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, -1),
                    "Helvetica-Bold",
                ),
                (
                    "ALIGN",
                    (1, 0),
                    (1, 0),
                    "RIGHT",
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.HexColor(
                        "#B9CEDD"
                    ),
                ),
            ]
        )
    )

    story.append(
        final_table
    )

    doc.build(story)

    output.seek(0)

    return output.getvalue()


# ============================================================
# HEADER
# ============================================================

st.html(
    """
    <div class="mdc-title-card">
        <div class="mdc-title">
            Eaton MDC Solution Configurator
        </div>

        <div class="mdc-subtitle">
            Modular Data Center Solution Configuration &amp; Pricing
        </div>
    </div>
    """
)


# ============================================================
# TOP CUSTOMER / USER INFO
# ============================================================

top_customer_col, top_info_col = st.columns(
    [2.2, 5.8],
    gap="small"
)

with top_customer_col:

    customer_name = st.text_input(
        "Customer Name",
        value=st.session_state.customer_name,
        key="customer_name_input",
        placeholder="Enter customer name",
        label_visibility="collapsed",
    )

    st.session_state.customer_name = (
        customer_name.strip()
    )

with top_info_col:

    user_info_placeholder = st.empty()


# ============================================================
# SIDEBAR ACCESS
# ============================================================

with st.sidebar:

    st.markdown(
        "### User Access"
    )

    selected_mode = st.radio(
        "Access",
        [
            "Sales",
            "Internal – MDC",
        ],
        index=(
            0
            if st.session_state.mode
            == "Sales"
            else 1
        ),
    )

    if selected_mode != st.session_state.mode:

        st.session_state.mode = selected_mode

        if selected_mode == "Sales":
            st.session_state.authenticated = False

        st.rerun()

    if st.session_state.mode == "Internal – MDC":

        password = st.text_input(
            "Internal Password",
            type="password"
        )

        if st.button(
            "Unlock",
            use_container_width=True
        ):

            if password == internal_password():

                st.session_state.authenticated = True

                st.success(
                    "Internal access enabled."
                )

            else:

                st.session_state.authenticated = False

                st.error(
                    "Invalid password."
                )

        if st.session_state.authenticated:

            if st.button(
                "Lock Internal Access",
                use_container_width=True
            ):

                st.session_state.authenticated = False
                st.rerun()


is_internal = (
    st.session_state.mode
    == "Internal – MDC"
    and st.session_state.authenticated
)


# ============================================================
# CUSTOMER DETAILS
# ============================================================

with st.expander(
    "Customer Details",
    expanded=False
):

    c1, c2 = st.columns(2)

    with c1:

        st.session_state.customer_place = (
            st.text_input(
                "Customer Place",
                value=st.session_state.customer_place,
            )
        )

    with c2:

        st.session_state.problem = (
            st.text_input(
                "Problem Description",
                value=st.session_state.problem,
            )
        )

    st.session_state.solution = (
        st.text_input(
            "Solution",
            value=st.session_state.solution,
        )
    )


# ============================================================
# MAIN TWO-COLUMN AREA
# ============================================================

main_left, main_right = st.columns(
    [1.0, 1.15],
    gap="medium"
)


# ============================================================
# LEFT PANEL
# ============================================================

with main_left:

    # ========================================================
    # 01. MDC TYPE & CONFIGURATION
    # ========================================================

    with st.container(border=True):

        st.markdown(
            '<div class="mdc-card-heading">'
            '<span class="mdc-number">01</span>'
            '<span>MDC TYPE &amp; CONFIGURATION</span>'
            '</div>',
            unsafe_allow_html=True,
        )

        mdc_type = st.radio(
            "MDC Type",
            [
                "Single Rack",
                "Multirack",
            ],
            horizontal=True,
            index=(
                0
                if st.session_state.mdc_type
                == "Single Rack"
                else 1
            ),
            key="mdc_type_selection",
        )

        if (
            mdc_type
            != st.session_state.mdc_type
        ):

            st.session_state.mdc_type = mdc_type

            st.session_state.configuration = (
                "Configuration 1"
            )

            st.session_state.accessory_qty = {}
            st.session_state.pdu_qty = {}

            st.session_state.pdu_type_selection = (
                "None"
            )

            st.session_state.pdu_model_selection = (
                None
            )

            st.session_state.configuration_id = (
                generate_configuration_id()
            )

            st.session_state.configuration_saved = (
                False
            )

            st.rerun()

        available = configs_df[
            configs_df["MDC Type"]
            == st.session_state.mdc_type
        ].copy()

        labels = (
            available["Configuration"]
            .tolist()
        )

        # ----------------------------------------------------
        # SINGLE RACK DISPLAY
        # ----------------------------------------------------

        single_display_names = {

            "Configuration 1":
                "Configuration 1 - "
                "1SR, 42U×800W×1200D, "
                "3.5KW, W/O Dehumidifier",

            "Configuration 2":
                "Configuration 2 - "
                "1SR, 42U×800W×1200D, "
                "3.5KW, Dehumidifier",

            "Configuration 3":
                "Configuration 3 - "
                "1SR, 42U×800W×1200D, "
                "7KW, W/O Dehumidifier",

            "Configuration 4":
                "Configuration 4 - "
                "1SR, 42U×800W×1200D, "
                "7KW, Dehumidifier",
        }

        # ----------------------------------------------------
        # STANDARD MULTIRACK DISPLAY
        # ----------------------------------------------------
        #
        # No rack customization.
        # Only standard configuration selection.
        #

        multirack_display_names = {

            f"Configuration {n}":
                f"Multirack Configuration {n}"
            for n in range(1, 10)
        }

        display_names = (
            single_display_names
            if st.session_state.mdc_type
            == "Single Rack"
            else multirack_display_names
        )

        if labels:

            selected_configuration = (
                st.selectbox(
                    "Select Configuration",
                    labels,

                    index=(
                        labels.index(
                            st.session_state.configuration
                        )
                        if st.session_state.configuration
                        in labels
                        else 0
                    ),

                    format_func=lambda x:
                        display_names.get(
                            x,
                            x
                        ),

                    key="configuration_selection",
                )
            )

            if (
                selected_configuration
                != st.session_state.configuration
            ):

                st.session_state.configuration = (
                    selected_configuration
                )

                st.session_state.accessory_qty = {}
                st.session_state.pdu_qty = {}

                st.session_state.pdu_type_selection = (
                    "None"
                )

                st.session_state.pdu_model_selection = (
                    None
                )

                st.session_state.configuration_id = (
                    generate_configuration_id()
                )

                st.session_state.configuration_saved = (
                    False
                )


    # ========================================================
    # 02. PDU SELECTION
    # ========================================================

    with st.container(border=True):

        st.markdown(
            '<div class="mdc-card-heading">'
            '<span class="mdc-number">02</span>'
            '<span>PDU SELECTION</span>'
            '</div>',
            unsafe_allow_html=True,
        )

        pdu_type_display = {
            "BASIC": "Basic PDU",
            "METERED": "Metered PDU",
            "SWITCHED": "Switched PDU",
            "MANAGED": "Managed PDU",
        }

        excel_pdu_types = []

        if not pdus_df.empty:

            excel_pdu_types = [
                x
                for x in (
                    pdus_df["Type"]
                    .astype(str)
                    .str.strip()
                    .str.upper()
                    .unique()
                    .tolist()
                )
                if x
            ]

        pdu_types = (
            ["None"]
            + [
                pdu_type_display.get(
                    x,
                    f"{x.title()} PDU"
                )
                for x in excel_pdu_types
            ]
        )

        pdu_col1, pdu_col2, pdu_col3 = (
            st.columns(
                [1.0, 2.0, 0.55],
                gap="small"
            )
        )

        with pdu_col1:

            previous_pdu_type = (
                st.session_state.get(
                    "pdu_type_selection",
                    "None"
                )
            )

            if (
                previous_pdu_type
                not in pdu_types
            ):
                previous_pdu_type = "None"

            selected_pdu_type = (
                st.selectbox(
                    "PDU Type",
                    pdu_types,
                    index=pdu_types.index(
                        previous_pdu_type
                    ),
                    key="pdu_type_selection",
                )
            )

        selected_part = None

        with pdu_col2:

            if selected_pdu_type != "None":

                reverse_type = {
                    v: k
                    for k, v
                    in pdu_type_display.items()
                }

                excel_pdu_type = (
                    reverse_type.get(
                        selected_pdu_type,
                        selected_pdu_type
                        .replace(
                            " PDU",
                            ""
                        )
                        .upper(),
                    )
                )

                filtered_pdus = pdus_df[
                    pdus_df["Type"]
                    .astype(str)
                    .str.strip()
                    .str.upper()
                    == excel_pdu_type
                ].copy()

                if not filtered_pdus.empty:

                    pdu_options = [
                        (
                            f'{clean_text(r["Part Code"])} '
                            f'— '
                            f'{clean_text(r["Description"])}'
                        )
                        for _, r
                        in filtered_pdus.iterrows()
                    ]

                    old_selection = (
                        st.session_state.get(
                            "pdu_model_selection"
                        )
                    )

                    pdu_index = (
                        pdu_options.index(
                            old_selection
                        )
                        if old_selection
                        in pdu_options
                        else 0
                    )

                    selected_pdu = (
                        st.selectbox(
                            "Select PDU",
                            pdu_options,
                            index=pdu_index,
                            key="pdu_model_selection",
                        )
                    )

                    selected_row = (
                        filtered_pdus.iloc[
                            pdu_options.index(
                                selected_pdu
                            )
                        ]
                    )

                    selected_part = (
                        clean_text(
                            selected_row[
                                "Part Code"
                            ]
                        )
                    )

                else:

                    st.warning(
                        f"No {selected_pdu_type} "
                        "options found."
                    )

        with pdu_col3:

            if selected_part:

                current_pdu_qty = int(
                    numeric(
                        st.session_state.pdu_qty.get(
                            selected_part,
                            1
                        )
                    )
                )

                pdu_quantity = (
                    st.number_input(
                        "Qty",
                        min_value=1,
                        max_value=999,
                        step=1,
                        value=current_pdu_qty,
                        key=(
                            f"pdu_quantity_"
                            f"{selected_part}"
                        ),
                    )
                )

                st.session_state.pdu_qty = {
                    selected_part:
                    pdu_quantity
                }

            else:

                st.number_input(
                    "Qty",
                    min_value=1,
                    max_value=999,
                    value=1,
                    step=1,
                    disabled=True,
                    key="pdu_quantity_none",
                )

                st.session_state.pdu_qty = {}


# ============================================================
# RIGHT PANEL
# ============================================================

with main_right:

    with st.container(border=True):

        st.markdown(
            '<div class="mdc-card-heading">'
            '<span class="mdc-number">03</span>'
            '<span>OTHER ACCESSORIES</span>'
            '</div>',
            unsafe_allow_html=True,
        )

        def excel_optional_rows(
            keyword=None
        ):

            if accessories_df.empty:
                return pd.DataFrame()

            if not keyword:
                return accessories_df.copy()

            key = str(keyword).upper()

            mask = (
                accessories_df[
                    "Part Code"
                ]
                .astype(str)
                .str.upper()
                .str.contains(
                    key,
                    na=False
                )
                |
                accessories_df[
                    "Description"
                ]
                .astype(str)
                .str.upper()
                .str.contains(
                    key,
                    na=False
                )
            )

            return accessories_df[
                mask
            ].copy()


        def remove_rows(rows):

            for _, r in rows.iterrows():

                part = clean_text(
                    r["Part Code"]
                )

                if part:
                    st.session_state.accessory_qty.pop(
                        part,
                        None
                    )


        def add_rows(
            rows,
            quantity=1
        ):

            for _, r in rows.iterrows():

                part = clean_text(
                    r["Part Code"]
                )

                if part:

                    st.session_state.accessory_qty[
                        part
                    ] = quantity


        # ----------------------------------------------------
        # FIRE SUPPRESSION
        # ----------------------------------------------------

        fire_rows = excel_optional_rows(
            "FIRE"
        )

        external_fire = fire_rows[
            fire_rows[
                "Description"
            ]
            .astype(str)
            .str.upper()
            .str.contains(
                "EXTERNAL",
                na=False
            )
        ].copy()

        internal_fire = fire_rows[
            ~fire_rows[
                "Description"
            ]
            .astype(str)
            .str.upper()
            .str.contains(
                "EXTERNAL",
                na=False
            )
        ].copy()

        fire_current = "None"

        if any(
            numeric(
                st.session_state.accessory_qty.get(
                    p,
                    0
                )
            ) > 0
            for p
            in external_fire[
                "Part Code"
            ]
            .astype(str)
            .str.strip()
        ):

            fire_current = "External"

        elif any(
            numeric(
                st.session_state.accessory_qty.get(
                    p,
                    0
                )
            ) > 0
            for p
            in internal_fire[
                "Part Code"
            ]
            .astype(str)
            .str.strip()
        ):

            fire_current = "Internal"


        fire_selection = st.radio(
            "Fire Suppression",
            [
                "None",
                "External",
                "Internal"
            ],
            index=[
                "None",
                "External",
                "Internal"
            ].index(
                fire_current
            ),
            horizontal=True,
            key="fire_suppression_selection",
        )

        remove_rows(
            external_fire
        )

        remove_rows(
            internal_fire
        )

        if fire_selection == "External":

            add_rows(
                external_fire,
                1
            )

        elif fire_selection == "Internal":

            add_rows(
                internal_fire,
                1
            )


        # ----------------------------------------------------
        # CAMERA
        # ----------------------------------------------------

        camera_rows = excel_optional_rows(
            "CAMERA"
        )

        camera_parts = (
            camera_rows[
                "Part Code"
            ]
            .astype(str)
            .str.strip()
            .tolist()
            if not camera_rows.empty
            else []
        )

        camera_current = any(
            numeric(
                st.session_state.accessory_qty.get(
                    p,
                    0
                )
            ) > 0
            for p in camera_parts
        )

        cam_col1, cam_col2 = st.columns(
            [4.5, 1.15],
            gap="small",
            vertical_alignment="center"
        )

        with cam_col1:

            camera_selection = (
                st.checkbox(
                    "Camera",
                    value=camera_current,
                    key="camera_system_selection",
                )
            )

        with cam_col2:

            if camera_selection:

                current_camera_qty = max(
                    [
                        int(
                            numeric(
                                st.session_state.accessory_qty.get(
                                    p,
                                    0
                                )
                            )
                        )
                        for p in camera_parts
                        if numeric(
                            st.session_state.accessory_qty.get(
                                p,
                                0
                            )
                        ) > 0
                    ]
                    + [1]
                )

                camera_qty = (
                    st.number_input(
                        "Qty",
                        min_value=1,
                        max_value=999,
                        value=int(
                            st.session_state.get(
                                "camera_quantity",
                                current_camera_qty
                            )
                        ),
                        step=1,
                        key="camera_quantity",
                    )
                )

            else:

                camera_qty = 0

        if camera_selection:

            add_rows(
                camera_rows,
                int(camera_qty)
            )

        else:

            remove_rows(
                camera_rows
            )


        # ----------------------------------------------------
        # OTHER ACCESSORIES
        # ----------------------------------------------------

        other_accessory_keywords = [
            (
                "KEYBOARD",
                "Rotating Keyboard Tray"
            ),
            (
                "CABLE MANAGER",
                "Cable Manager"
            ),
            (
                "TOP CABLE TRAY",
                "Top Cable Tray"
            ),
            (
                "BRUSH PANEL",
                "Brush Panel"
            ),
        ]

        for keyword, fallback_label in (
            other_accessory_keywords
        ):

            rows = excel_optional_rows(
                keyword
            )

            if rows.empty:
                continue

            for _, r in rows.iterrows():

                part = clean_text(
                    r["Part Code"]
                )

                description = clean_text(
                    r["Description"]
                )

                if not part:
                    continue

                current_qty = int(
                    numeric(
                        st.session_state.accessory_qty.get(
                            part,
                            0
                        )
                    )
                )

                acc_col1, acc_col2 = (
                    st.columns(
                        [4.5, 1.15],
                        gap="small",
                        vertical_alignment="center"
                    )
                )

                with acc_col1:

                    selected = st.checkbox(
                        (
                            description
                            if description
                            else fallback_label
                        ),
                        value=(
                            current_qty > 0
                        ),
                        key=(
                            f"other_acc_"
                            f"{part}"
                        ),
                    )

                with acc_col2:

                    if selected:

                        qty = st.number_input(
                            "Qty",
                            min_value=1,
                            max_value=999,
                            step=1,
                            value=(
                                current_qty
                                if current_qty > 0
                                else 1
                            ),
                            key=(
                                f"other_qty_"
                                f"{part}"
                            ),
                            label_visibility="collapsed",
                        )

                        st.session_state.accessory_qty[
                            part
                        ] = qty

                    else:

                        st.session_state.accessory_qty.pop(
                            part,
                            None
                        )


# ============================================================
# USER INFORMATION
# ============================================================

st.session_state.user_code = (
    generate_user_code()
)

render_user_info_panel(
    user_info_placeholder
)


# ============================================================
# 05. FINAL BOQ
# ============================================================

bom = build_bom()

(
    base_cost,
    optional_cost,
    pdu_cost,
    total_cost,
) = cost_summary(bom)

margin_pct = float(
    st.session_state.margin_pct
)

freight = float(
    st.session_state.freight
)

installation = float(
    st.session_state.installation
)

if not bom.empty:

    (
        bom_with_price,
        margin_price,
        final_selling_price,
    ) = add_selling_prices(
        bom,
        total_cost,
        margin_pct,
        freight,
        installation,
    )

else:

    bom_with_price = bom.copy()

    margin_price = (
        total_cost
        / (
            1
            - margin_pct / 100
        )
        if margin_pct < 100
        else 0.0
    )

    final_selling_price = (
        margin_price
        + freight
        + installation
    )


st.html(
    f"""
    <div style="
        display:flex;
        justify-content:space-between;
        align-items:center;
        gap:15px;
        background:#003B71;
        color:white;
        padding:7px 12px;
        border-radius:5px;
        margin:10px 0 8px 0;
    ">

        <div style="
            font-size:14px;
            font-weight:700;
        ">
            5. FINAL BOQ
        </div>

        <div style="
            display:flex;
            align-items:center;
            gap:8px;
            white-space:nowrap;
        ">

            <span style="
                font-size:11px;
                font-weight:700;
            ">
                FINAL SELLING PRICE
            </span>

            <span style="
                font-size:16px;
                font-weight:700;
            ">
                {money(final_selling_price)}
            </span>

        </div>
    </div>
    """
)


# ============================================================
# FINAL BOQ TABLE
# ============================================================

if not bom.empty:

    structure = bom_with_price[
        [
            "S.No.",
            "Part Code",
            "Description",
            "Quantity",
            "UOM",
            "Unit Price",
            "Total Price",
        ]
    ].copy()

    # --------------------------------------------------------
    # SERIAL NUMBERING
    # --------------------------------------------------------

    selected_config_components = (
        selected_components()
    )

    cooling_part_codes = set()

    if not selected_config_components.empty:

        cooling_rows = (
            selected_config_components.tail(3)
        )

        cooling_part_codes = set(
            cooling_rows[
                "Part Code"
            ]
            .dropna()
            .astype(str)
            .str.strip()
        )

    new_serial = []

    mdc_sub_no = 0
    cooling_sub_no = 16
    pdu_sub_no = 0
    accessory_sub_no = 0

    for idx, row in structure.iterrows():

        part_code = clean_text(
            row["Part Code"]
        )

        component_type = clean_text(
            bom.loc[
                row.name,
                "Component Type"
            ]
        )

        if part_code == "801029209":

            new_serial.append("1")
            continue

        if (
            component_type
            == "Base (Configuration)"
        ):

            if part_code in cooling_part_codes:

                cooling_sub_no += 1

                new_serial.append(
                    f"1.{cooling_sub_no}"
                )

            else:

                mdc_sub_no += 1

                new_serial.append(
                    f"1.{mdc_sub_no}"
                )

            continue

        if component_type == "PDU":

            pdu_sub_no += 1

            new_serial.append(
                f"2.{pdu_sub_no}"
            )

            continue

        if (
            component_type
            == "Optional Accessory"
        ):

            accessory_sub_no += 1

            new_serial.append(
                f"3.{accessory_sub_no}"
            )

            continue

        new_serial.append("")

    structure["New S.No."] = new_serial


    # --------------------------------------------------------
    # HTML
    # --------------------------------------------------------

    html = """
    <style>

    .final-structure-wrap {
        width:100%;
        overflow-x:auto;
    }

    .final-structure-table {
        width:100%;
        border-collapse:collapse;
        table-layout:fixed;
        font-family:Arial,sans-serif;
        font-size:12px;
        border:1px solid #D9E1E8;
    }

    .final-structure-table th {
        background:#F4F6F8;
        color:#555;
        font-weight:600;
        text-align:left;
        padding:7px 7px;
        border-bottom:1px solid #D9E1E8;
    }

    .final-structure-table td {
        padding:6px 7px;
        border-bottom:1px solid #E5E7EB;
        color:#333;
        vertical-align:middle;
        overflow-wrap:anywhere;
    }

    .main-mdc-row td {
        background:#003B71;
        color:white !important;
        font-weight:700;
        font-size:13px;
        text-align:center !important;
        padding:8px 7px;
    }

    .section-heading td {
        background:#005EB8;
        color:white !important;
        font-weight:700;
        font-size:12px;
        text-align:center !important;
        padding:7px 8px;
    }

    .selected-config-title-row td {
        background:#D9EAF7;
        color:#003B71 !important;
        font-weight:700;
        font-size:12px;
        text-align:center !important;
        padding:7px 8px;
    }

    .serial {
        width:7%;
        text-align:center !important;
    }

    .part-code {
        width:13%;
    }

    .description {
        width:40%;
    }

    .quantity {
        width:8%;
        text-align:center !important;
    }

    .uom {
        width:7%;
        text-align:center !important;
    }

    .unit-price {
        width:12%;
        text-align:right !important;
        white-space:nowrap;
    }

    .total-price {
        width:13%;
        text-align:right !important;
        white-space:nowrap;
    }

    </style>

    <div class="final-structure-wrap">

    <table class="final-structure-table">

        <thead>
            <tr>
                <th class="serial">S.No.</th>
                <th class="part-code">Part Code</th>
                <th class="description">Description</th>
                <th class="quantity">Qty</th>
                <th class="uom">UOM</th>
                <th class="unit-price">Unit Price</th>
                <th class="total-price">Total Price</th>
            </tr>
        </thead>

        <tbody>
    """

    selected_config = (
        selected_config_record()
    )

    selected_config_title = (
        clean_text(
            selected_config.get(
                "Configuration Title",
                ""
            )
        )
        if selected_config is not None
        else ""
    )

    if selected_config_title:

        html += f"""
        <tr class="selected-config-title-row">
            <td colspan="7">
                {selected_config_title}
            </td>
        </tr>
        """

    cooling_heading_added = False
    accessories_heading_added = False
    pdu_heading_added = False

    for _, row in structure.iterrows():

        part_code = clean_text(
            row["Part Code"]
        )

        description = clean_text(
            row["Description"]
        )

        quantity_value = numeric(
            row["Quantity"]
        )

        quantity = (
            f"{quantity_value:g}"
        )

        uom = clean_text(
            row["UOM"]
        )

        serial_no = clean_text(
            row["New S.No."]
        )

        component_type = clean_text(
            bom.loc[
                row.name,
                "Component Type"
            ]
        )

        unit_price = numeric(
            row["Unit Price"]
        )

        total_price = numeric(
            row["Total Price"]
        )

        if (
            serial_no == ""
            and (
                "SINGLE RACK MDC"
                in description.upper()
                or
                "MULTIRACK MDC"
                in description.upper()
            )
        ):

            html += f"""
            <tr class="main-mdc-row">
                <td colspan="7">
                    {description}
                </td>
            </tr>
            """

            continue

        if (
            part_code
            in cooling_part_codes
            and not cooling_heading_added
        ):

            html += """
            <tr class="section-heading">
                <td colspan="7">
                    COOLING UNIT
                </td>
            </tr>
            """

            cooling_heading_added = True

        if (
            component_type
            == "Optional Accessory"
            and not accessories_heading_added
        ):

            html += """
            <tr class="section-heading">
                <td colspan="7">
                    OTHER ACCESSORIES
                </td>
            </tr>
            """

            accessories_heading_added = True

        if (
            component_type == "PDU"
            and not pdu_heading_added
        ):

            html += """
            <tr class="section-heading">
                <td colspan="7">
                    PDU
                </td>
            </tr>
            """

            pdu_heading_added = True

        html += f"""
        <tr>

            <td class="serial">
                {serial_no}
            </td>

            <td class="part-code">
                {part_code}
            </td>

            <td class="description">
                {description}
            </td>

            <td class="quantity">
                {quantity}
            </td>

            <td class="uom">
                {uom}
            </td>

            <td class="unit-price">
                {money(unit_price)}
            </td>

            <td class="total-price">
                {money(total_price)}
            </td>

        </tr>
        """

    html += f"""
        <tr>

            <td colspan="6"
                style="
                    text-align:right;
                    font-weight:700;
                    padding:8px 7px;
                    background:#F7FBFF;
                    color:#003B71;
                "
            >
                FINAL SELLING PRICE
            </td>

            <td class="total-price"
                style="
                    font-weight:700;
                    background:#F7FBFF;
                    color:#003B71;
                "
            >
                {money(final_selling_price)}
            </td>

        </tr>

        </tbody>

    </table>

    </div>
    """

    st.html(html)

else:

    st.info(
        "No components selected for "
        "the current configuration."
    )


# ============================================================
# 06. INTERNAL COST SUMMARY
# ============================================================

if is_internal:

    section_header(
        "6. COST SUMMARY — INTERNAL ONLY"
    )

    a, b, c, d = st.columns(4)

    with a:
        price_box(
            "Base Cost",
            base_cost
        )

    with b:
        price_box(
            "Optional Cost",
            optional_cost
        )

    with c:
        price_box(
            "PDU Cost",
            pdu_cost
        )

    with d:
        price_box(
            "Total Cost",
            total_cost
        )


# ============================================================
# 07. COST TO SELLING PRICE
# ============================================================

if is_internal:

    section_header(
        "7. COST TO SELLING PRICE — INTERNAL ONLY"
    )

    p1, p2, p3, p4 = st.columns(4)

    with p1:

        st.session_state.margin_pct = (
            st.number_input(
                "Margin (%)",
                min_value=0.0,
                max_value=99.0,
                value=float(
                    st.session_state.margin_pct
                ),
                step=0.5,
            )
        )

    with p2:

        st.session_state.freight = (
            st.number_input(
                "Freight",
                min_value=0.0,
                value=float(
                    st.session_state.freight
                ),
                step=500.0,
            )
        )

    with p3:

        st.session_state.installation = (
            st.number_input(
                "Installation",
                min_value=0.0,
                value=float(
                    st.session_state.installation
                ),
                step=500.0,
            )
        )

    with p4:

        st.session_state.warranty_pct = (
            st.number_input(
                "Warranty (%)",
                min_value=0.0,
                max_value=100.0,
                value=float(
                    st.session_state.warranty_pct
                ),
                step=0.5,
            )
        )

    margin_pct = float(
        st.session_state.margin_pct
    )

    freight = float(
        st.session_state.freight
    )

    installation = float(
        st.session_state.installation
    )

    warranty_pct = float(
        st.session_state.warranty_pct
    )

    margin_price = (
        total_cost
        / (
            1
            - margin_pct / 100
        )
        if margin_pct < 100
        else 0.0
    )

    final_selling_price = (
        margin_price
        + freight
        + installation
    )

    warranty_amount = (
        margin_price
        * warranty_pct
        / 100
    )

    a, b, c, d = st.columns(4)

    with a:
        price_box(
            "Margin Price",
            margin_price
        )

    with b:
        price_box(
            "After Freight",
            margin_price + freight
        )

    with c:
        price_box(
            "Final Selling Price",
            final_selling_price
        )

    with d:
        price_box(
            "Warranty Amount",
            warranty_amount
        )


# ============================================================
# 08. DOWNLOADS
# ============================================================

section_header(
    "8. DOWNLOADS"
)

if not bom.empty:

    internal_cost_data = [
        ["Base Cost", base_cost],
        ["Optional Cost", optional_cost],
        ["PDU Cost", pdu_cost],
        ["Total Cost", total_cost],
        ["Margin %", margin_pct],
        ["Margin Price", margin_price],
        ["Freight", freight],
        ["Installation", installation],
        ["Final Selling Price", final_selling_price],
    ]

    sales_excel = excel_bytes(
        internal=False,
        bom=bom_with_price,
        final_price=final_selling_price,
    )

    sales_pdf = pdf_bytes(
        internal=False,
        bom=bom_with_price,
        final_price=final_selling_price,
    )

    if is_internal:

        internal_excel = excel_bytes(
            internal=True,
            bom=bom_with_price,
            final_price=final_selling_price,
            cost_data=internal_cost_data,
        )

        internal_pdf = pdf_bytes(
            internal=True,
            bom=bom_with_price,
            final_price=final_selling_price,
        )

        st.markdown(
            "**Internal – MDC**"
        )

        i1, i2 = st.columns(2)

        with i1:

            st.download_button(
                "⬇️ Download Internal Excel",
                data=internal_excel,
                file_name=(
                    "MDC_Internal_Cost.xlsx"
                ),
                mime=(
                    "application/vnd.openxmlformats-"
                    "officedocument.spreadsheetml.sheet"
                ),
                use_container_width=True,
                on_click=handle_excel_download,
            )

        with i2:

            st.download_button(
                "📄 Download Internal PDF",
                data=internal_pdf,
                file_name=(
                    "MDC_Internal_Cost.pdf"
                ),
                mime="application/pdf",
                use_container_width=True,
                on_click=handle_excel_download,
            )

        st.markdown(
            "**Sales**"
        )

    s1, s2 = st.columns(2)

    with s1:

        st.download_button(
            "⬇️ Download Sales Excel",
            data=sales_excel,
            file_name=(
                "MDC_Sales_Output.xlsx"
            ),
            mime=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
            use_container_width=True,
            on_click=handle_excel_download,
        )

    with s2:

        st.download_button(
            "📄 Download Sales PDF",
            data=sales_pdf,
            file_name=(
                "MDC_Sales_Output.pdf"
            ),
            mime="application/pdf",
            use_container_width=True,
            on_click=handle_excel_download,
        )

else:

    st.info(
        "Select a configuration with "
        "available BOM data before downloading."
    )


# ============================================================
# 10. CONFIGURATION HISTORY
# INTERNAL ONLY
# ============================================================

if is_internal:

    section_header(
        "10. CONFIGURATION HISTORY"
    )

    conn = sqlite3.connect(
        TRACKING_DB
    )

    history_df = pd.read_sql_query(
        """
        SELECT
            user_code AS "User Code",
            customer_name AS "Customer Name"
        FROM user_history
        ORDER BY id DESC
        """,
        conn,
    )

    conn.close()

    if not history_df.empty:

        st.dataframe(
            history_df,
            use_container_width=True,
            hide_index=True,
            height=220,
        )

    else:

        st.info(
            "No user history available yet."
        )


# ============================================================
# FOOTER
# ============================================================

st.divider()

st.caption(
    "Eaton MDC Solution Configurator"
)
